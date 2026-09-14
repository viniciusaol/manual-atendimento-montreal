import openpyxl
import json

path = r"C:\Users\vinic\Downloads\Politica%20precificacao%20-%20Matricula%20proporcional.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print("Sheets in workbook:", wb.sheetnames)

all_data = {}

for name in wb.sheetnames:
    sheet = wb[name]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        # Ignore completely empty rows
        if any(cell is not None for cell in row):
            clean_row = [str(c).strip() if c is not None else "" for c in row]
            rows.append(clean_row)
    all_data[name] = rows
    print(f"\n=== SHEET: {name} ({len(rows)} rows) ===")
    for r in rows[:15]:
        print("  |  ".join(r[:8]))

with open("data/pricing_excel_dump.json", "w", encoding="utf-8") as f:
    json.dump(all_data, f, ensure_ascii=False, indent=2)

print("\nSaved full dump to data/pricing_excel_dump.json")
