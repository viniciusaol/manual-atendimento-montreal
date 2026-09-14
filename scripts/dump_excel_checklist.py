import openpyxl
import json

path = r'C:\Users\vinic\Downloads\Checklist_Diario_Montreal_Tenis_Clube (1).xlsx'
wb = openpyxl.load_workbook(path)

data = {}

for sheet in wb.sheetnames:
    ws = wb[sheet]
    rows = []
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rows.append(row_vals)
    data[sheet] = rows

with open('data/checklist_excel_dump.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, ensure_ascii=False, default=str)

print("Dumped sheets:", list(data.keys()))
